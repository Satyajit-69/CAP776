import openpyxl
import math

path = r"C:\Users\satya\Desktop\datasheet\12600637.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)

print(wb.sheetnames)

ws = wb["Daily Log"]

def get_values(data,start, end, col):
    

    for row in ws[f"{col}{start}:{col}{end}"]:
        for cell in row:

            if cell.value is None:
                continue

            if cell.value == 0:
                continue

            data.append(cell.value)

   

#SRI
def calculate_avg_sleep():
    sleep_data = []
    get_values(sleep_data, 7, 53, "B")
    sri = sum(sleep_data) / len(sleep_data)
    print(f"my SRI : {sri:.2f} minutes")
    return sri


# other activity 
def other_activity():
    other_data = []
    get_values(other_data, 7, 53, "H")
    otherA = sum(other_data) / len(other_data)
    print(f"other datas : {otherA:.2f} minutes")

other_activity()



# free time 
def free_time():
    free_time = []
    get_values(free_time, 7, 53, "J")   
    freeT = sum(free_time) / len(free_time)
    print(f"avg free times : {freeT:.2f} minutes")

free_time()




#tpi - total productivity index
def total_productivity():

    coding_data = []

    get_values(coding_data, 7, 53, "E")

    tpi = sum(coding_data) / len(coding_data)

    print(f"my TPI : {tpi:.2f} minutes")

    return tpi


#AAI - academic activity index
def academic_activity():

    study_data = []
    class_data = []

    get_values(study_data, 7, 53, "D")
    print("study_time : " , sum(study_data) / len(study_data) )

    get_values(class_data, 7, 53, "F")
    print("class data  : " , sum(class_data) / len(class_data) )
    

    academic_data = []

    for i in range(len(study_data)):
        academic_data.append(study_data[i])

    for i in range(len(class_data)):
        academic_data.append(class_data[i])

    aai = sum(academic_data) / len(academic_data)

    print(f"my AAI : {aai:.2f} minutes")

    return aai


#PhAI -Physical Activity Index

def physical_activity():

    fitness_data = []
    get_values(fitness_data, 7, 53, "C")
    phai = sum(fitness_data) / len(fitness_data)
    print(f"PhAI : {phai:.2f} minutes")

    return phai



# ABI -Activity balance Index


def activity_balance():

    total_time_data = []
    get_values(total_time_data, 7, 53, "I")
    free_time_data = []

    for value in total_time_data:
        free_time_data.append(1440 - value)

    abi = sum(free_time_data) / len(free_time_data)

    print(f"ABI : {abi:.2f} minutes")

    return abi


# TUI -Time utilization Index
def time_utilization():

    total_time_data = []
    get_values(total_time_data, 7, 53, "I")
    tui = sum(total_time_data) / len(total_time_data)

    print(f"TUI : {tui:.2f} minutes")

    return tui


# DCI - Data Continuity Index
def data_continuity():

    sleep_data = []

    get_values(sleep_data, 9, 52, "B")

    expected_days = 44

    dci = (len(sleep_data) / expected_days) * 100

    print(f"DCI : {dci:.2f} %")

    return dci


# EI - Experience Index
def experience_index():

    feeling = {
        "excellent": 5,
        "very good": 5,
        "good": 4,
        "normal": 3,
        "neutral": 3,
        "low": 2,
        "bad": 1,
        # "stressed" : 1 ,
    }

    satisfaction_rate = {
        "very satisfied": 5,
        "satisfied": 4,
        "satishfied": 4,
        "normal": 3,
        "neutral": 3,
        "unsatisfied": 2,
        "very unsatisfied": 1,
        
    }

    energy = {
        "high": 3,
        "medium": 2,
        "low": 1
    }

    experience_data = []

    for r in range(7, 53):

        if ws[f"B{r}"].value is None:
            continue

        #extracting the cell value
        feeling_value = ws[f"K{r}"].value
        satisfaction_value = ws[f"L{r}"].value
        energy_value = ws[f"M{r}"].value

        if feeling_value is None or satisfaction_value is None:
            continue #invalid
        
        #getting the scores based on the values
        feeling_score = feeling[feeling_value.strip().lower()]
        satisfaction_score = satisfaction_rate[
            satisfaction_value.strip().lower()
        ]

        #no energy value case 
        if energy_value is None:
            score = (feeling_score + satisfaction_score) / 2
        else:
            energy_score = energy[
                energy_value.strip().lower()
            ]
            score = (feeling_score + satisfaction_score + energy_score) / 3

        experience_data.append(score)

    ei = sum(experience_data) / len(experience_data)

    print(f"EI : {ei:.2f} / 5")

    return ei



sri = calculate_avg_sleep()
tpi = total_productivity()
aai = academic_activity()
phai = physical_activity()
abi = activity_balance()
tui = time_utilization()
dci = data_continuity()
ei = experience_index()


#calculating personal activity index
pai = (0.15 * tpi) + (0.20 * aai) + (0.15 * phai) + (0.20 * sri) + (0.15 * tui) + (0.10 * ei) + (0.05 * dci)
print(f"personal activity index  : {pai:.2f}")



def correlation(x, y):

    n = len(x)

    avg_of_x = sum(x) / n
    avg_of_y = sum(y) / n

    numerator = 0
    x_difference = 0
    y_difference = 0

    for i in range(n):

        numerator += ((x[i] - avg_of_x) * (y[i] - avg_of_y))

        x_difference +=  (x[i] - avg_of_x) ** 2
        y_difference += (y[i] - avg_of_y) ** 2

    denominator = (x_difference * y_difference) ** 0.5

    if denominator == 0:
        return 0

    return numerator / denominator


def relationship_analysis():

    coding_data = []
    coding_energy = []

    sleep_data = []
    sleep_energy = []

    study_data = []
    study_satisfaction = []

    energy = {
        "high": 3,
        "medium": 2,
        "low": 1
    }

    satisfaction = {
        "very satisfied": 5,
        "satisfied": 4,
        "satishfied": 4,
        "normal": 3,
        "neutral": 3,
        "unsatisfied": 2,
        "very unsatisfied": 1
    }


    for r in range(7, 54):

        coding = ws[f"E{r}"].value
        sleep = ws[f"B{r}"].value
        study = ws[f"D{r}"].value

        energy_value = ws[f"M{r}"].value
        satisfaction_value = ws[f"L{r}"].value


        # Coding -> Energy
        if coding is not None and energy_value is not None:

            energy_score = energy[
                energy_value.strip().lower()
            ]

            coding_data.append(coding)
            coding_energy.append(energy_score)


        # Sleep -> Energy
        if sleep is not None and energy_value is not None:

            energy_score = energy[
                energy_value.strip().lower()
            ]

            sleep_data.append(sleep)
            sleep_energy.append(energy_score)


        # Study -> Satisfaction
        if study is not None and satisfaction_value is not None:

            satisfaction_score = satisfaction[
                satisfaction_value.strip().lower()
            ]

            study_data.append(study)
            study_satisfaction.append(satisfaction_score)


    print()
    print("--RELATIONSHIP ANALYSIS ----------")

    coding_corr = correlation(coding_data, coding_energy)
    sleep_corr = correlation(sleep_data, sleep_energy)
    study_corr = correlation(study_data, study_satisfaction)

    print(f"Coding <-> Energy        : {coding_corr:.2f}")
    print(f"Sleep <-> Energy         : {sleep_corr:.2f}")
    print(f"Study <-> Satisfaction   : {study_corr:.2f}")




relationship_analysis() 




    